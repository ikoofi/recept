"""Локальная книга рецептов: Flask + MySQL, экспорт рецепта в DOCX."""
import io
import os
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pymysql
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from flask import Flask, abort, flash, redirect, render_template, request, send_file, url_for

BASE_DIR = Path(__file__).resolve().parent

# Порядок важен: так единицы показываются в выпадающем списке.
UNITS = {
    "g": "г",
    "kg": "кг",
    "ml": "мл",
    "l": "л",
    "tbsp": "ст. л.",
    "tsp": "ч. л.",
    "pcs": "шт.",
    "cloves": "зубч.",
    "to_taste": "по вкусу",
}


def load_env(path):
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip())


load_env(BASE_DIR / ".env")

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "local-recipe-book")


def get_db():
    return pymysql.connect(
        host=os.environ.get("DB_HOST", "localhost"),
        port=int(os.environ.get("DB_PORT", 3306)),
        user=os.environ.get("DB_USER", "recept"),
        password=os.environ.get("DB_PASSWORD", ""),
        database=os.environ.get("DB_NAME", "recept"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=False,
    )


def format_quantity(quantity, unit):
    if unit == "to_taste":
        return UNITS[unit]
    # Округление до сотых нужно для пересчитанных значений (например, 1/3 порции).
    q = Decimal(quantity).quantize(Decimal("0.01")).normalize()
    text = format(q, "f").replace(".", ",")
    return f"{text} {UNITS[unit]}"


app.jinja_env.filters["qty"] = lambda ing: format_quantity(ing["quantity"], ing["unit"])
app.jinja_env.filters["qty_input"] = lambda ing: (
    "" if ing["unit"] == "to_taste"
    else format(Decimal(ing["quantity"]).quantize(Decimal("0.01")).normalize(), "f")
)


def parse_main_qty(value):
    """Желаемое количество главного ингредиента из ?main_qty=, либо None."""
    try:
        q = Decimal((value or "").strip().replace(",", "."))
    except InvalidOperation:
        return None
    return q if q > 0 else None


def scale_recipe(recipe, main_qty):
    """Пересчитывает все количества пропорционально главному ингредиенту."""
    main = recipe["main"]
    if main_qty is None or main is None or not main["quantity"]:
        return recipe
    factor = main_qty / Decimal(main["quantity"])
    for ing in recipe["ingredients"]:
        ing["quantity"] = Decimal(ing["quantity"]) * factor
    return recipe


def fetch_recipe(recipe_id):
    with get_db() as conn, conn.cursor() as cur:
        cur.execute("SELECT * FROM recipes WHERE id = %s", (recipe_id,))
        recipe = cur.fetchone()
        if recipe is None:
            abort(404)
        cur.execute(
            "SELECT * FROM ingredients WHERE recipe_id = %s ORDER BY position", (recipe_id,)
        )
        recipe["ingredients"] = cur.fetchall()
        cur.execute("SELECT * FROM steps WHERE recipe_id = %s ORDER BY position", (recipe_id,))
        recipe["steps"] = cur.fetchall()
    recipe["main"] = next((i for i in recipe["ingredients"] if i["is_main"]), None)
    return recipe


def parse_form(form):
    """Собирает рецепт из формы. Возвращает (recipe, errors)."""
    errors = []
    title = form.get("title", "").strip()
    if not title:
        errors.append("Укажите название рецепта.")

    main_key = form.get("main")
    ingredients = []
    rows = zip(
        form.getlist("ing_key"),
        form.getlist("ing_name"),
        form.getlist("ing_qty"),
        form.getlist("ing_unit"),
    )
    for key, name, qty, unit in rows:
        name = name.strip()
        qty = qty.strip().replace(",", ".")
        if not name:
            continue
        if unit not in UNITS:
            errors.append(f"Неизвестная единица измерения у «{name}».")
            continue
        if unit == "to_taste":
            quantity = Decimal(0)
        else:
            try:
                quantity = Decimal(qty)
            except InvalidOperation:
                errors.append(f"Укажите количество для «{name}».")
                continue
            if quantity <= 0:
                errors.append(f"Количество для «{name}» должно быть больше нуля.")
                continue
        ingredients.append(
            {"name": name, "quantity": quantity, "unit": unit, "is_main": key == main_key}
        )

    if not ingredients:
        errors.append("Добавьте хотя бы один ингредиент.")
    elif sum(i["is_main"] for i in ingredients) != 1:
        errors.append("Отметьте главный ингредиент.")
    elif next(i for i in ingredients if i["is_main"])["unit"] == "to_taste":
        errors.append("У главного ингредиента должно быть количество, а не «по вкусу».")

    steps = [{"text": s.strip()} for s in form.getlist("step") if s.strip()]
    if not steps:
        errors.append("Добавьте хотя бы один шаг приготовления.")

    return {"title": title, "ingredients": ingredients, "steps": steps}, errors


def save_children(cur, recipe_id, recipe):
    cur.execute("DELETE FROM ingredients WHERE recipe_id = %s", (recipe_id,))
    cur.execute("DELETE FROM steps WHERE recipe_id = %s", (recipe_id,))
    cur.executemany(
        "INSERT INTO ingredients (recipe_id, position, name, quantity, unit, is_main)"
        " VALUES (%s, %s, %s, %s, %s, %s)",
        [
            (recipe_id, pos, i["name"], i["quantity"], i["unit"], i["is_main"])
            for pos, i in enumerate(recipe["ingredients"], 1)
        ],
    )
    cur.executemany(
        "INSERT INTO steps (recipe_id, position, text) VALUES (%s, %s, %s)",
        [(recipe_id, pos, s["text"]) for pos, s in enumerate(recipe["steps"], 1)],
    )


@app.route("/")
def index():
    q = request.args.get("q", "").strip()
    sql = (
        "SELECT r.id, r.title, i.name AS main_ingredient"
        " FROM recipes r LEFT JOIN ingredients i ON i.recipe_id = r.id AND i.is_main"
    )
    params = ()
    if q:
        sql += " WHERE r.title LIKE %s OR i.name LIKE %s"
        params = (f"%{q}%", f"%{q}%")
    sql += " ORDER BY r.title"
    with get_db() as conn, conn.cursor() as cur:
        cur.execute(sql, params)
        recipes = cur.fetchall()
    return render_template("index.html", recipes=recipes, q=q)


@app.route("/recipe/new", methods=["GET", "POST"])
def recipe_new():
    if request.method == "POST":
        recipe, errors = parse_form(request.form)
        if not errors:
            with get_db() as conn:
                with conn.cursor() as cur:
                    cur.execute("INSERT INTO recipes (title) VALUES (%s)", (recipe["title"],))
                    recipe_id = cur.lastrowid
                    save_children(cur, recipe_id, recipe)
                conn.commit()
            flash("Рецепт сохранён.")
            return redirect(url_for("recipe_view", recipe_id=recipe_id))
        for e in errors:
            flash(e, "error")
        return render_template("form.html", recipe=recipe, units=UNITS, mode="new")
    return render_template("form.html", recipe=None, units=UNITS, mode="new")


@app.route("/recipe/<int:recipe_id>")
def recipe_view(recipe_id):
    main_qty = parse_main_qty(request.args.get("main_qty"))
    recipe = scale_recipe(fetch_recipe(recipe_id), main_qty)
    return render_template("view.html", recipe=recipe, main_qty=main_qty, units=UNITS)


@app.route("/recipe/<int:recipe_id>/edit", methods=["GET", "POST"])
def recipe_edit(recipe_id):
    existing = fetch_recipe(recipe_id)
    if request.method == "POST":
        recipe, errors = parse_form(request.form)
        if not errors:
            with get_db() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        "UPDATE recipes SET title = %s WHERE id = %s", (recipe["title"], recipe_id)
                    )
                    save_children(cur, recipe_id, recipe)
                conn.commit()
            flash("Изменения сохранены.")
            return redirect(url_for("recipe_view", recipe_id=recipe_id))
        for e in errors:
            flash(e, "error")
        recipe["id"] = recipe_id
        return render_template("form.html", recipe=recipe, units=UNITS, mode="edit")
    return render_template("form.html", recipe=existing, units=UNITS, mode="edit")


@app.route("/recipe/<int:recipe_id>/delete", methods=["POST"])
def recipe_delete(recipe_id):
    with get_db() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM recipes WHERE id = %s", (recipe_id,))
        conn.commit()
    flash("Рецепт удалён.")
    return redirect(url_for("index"))


@app.route("/recipe/<int:recipe_id>/docx")
def recipe_docx(recipe_id):
    recipe = scale_recipe(fetch_recipe(recipe_id), parse_main_qty(request.args.get("main_qty")))
    doc = build_docx(recipe)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = re.sub(r'[\\/:*?"<>|]+', "_", recipe["title"]).strip() or "recipe"
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{filename}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def build_docx(recipe):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)

    title = doc.add_heading(recipe["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if recipe["main"]:
        p = doc.add_paragraph()
        p.add_run("Главный ингредиент: ").italic = True
        p.add_run(recipe["main"]["name"]).bold = True

    doc.add_heading("Ингредиенты", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for ing in recipe["ingredients"]:
        cells = table.add_row().cells
        name_run = cells[0].paragraphs[0].add_run(ing["name"])
        qty_run = cells[1].paragraphs[0].add_run(
            format_quantity(ing["quantity"], ing["unit"])
        )
        if ing["is_main"]:
            name_run.bold = qty_run.bold = True
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_heading("Приготовление", level=1)
    for step in recipe["steps"]:
        doc.add_paragraph(step["text"], style="List Number")

    return doc


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5000)), debug=True)
